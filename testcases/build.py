# -*- coding: utf-8 -*-
"""由 cases_data.py 生成 Markdown 测试用例文档与 Excel 版本。"""

import datetime
from collections import Counter, OrderedDict

from cases_data import CASES, FP, POINTS, SEVERITY, A_RULE, A_SPECIAL, A_PERM

TODAY = "2026-10-19"

FP_SECTION = {
    "FP-01": (A_RULE, "三种规则类型"),
    "FP-02": (A_RULE, "打卡人员"),
    "FP-03": (A_RULE, "三种打卡方式"),
    "FP-04": (A_RULE, "打卡地点与打卡范围"),
    "FP-05": (A_RULE, "上下班时间与容差"),
    "FP-06": (A_RULE, "外出打卡"),
    "FP-07": (A_SPECIAL, "拍照打卡"),
    "FP-08": (A_SPECIAL, "补卡申请"),
    "FP-09": (A_SPECIAL, "特殊日期打卡"),
    "FP-10": (A_SPECIAL, "打卡提醒"),
    "FP-11": (A_SPECIAL, "打卡应用的可见范围"),
    "FP-12": (A_SPECIAL, "打卡与审批的关联"),
    "FP-13": (A_PERM, "两类管理员"),
    "FP-14": (A_PERM, "查看与导出打卡记录"),
    "FP-15": (A_PERM, "API 导出数据"),
}

POINT_STATUS = {
    "P1-1": "覆盖", "P1-2": "覆盖", "P1-3": "覆盖", "P1-4": "覆盖",
    "P1-5": "覆盖（含待确认分支）", "P1-6": "覆盖",
    "P2-1": "覆盖", "P2-2": "覆盖", "P2-3": "覆盖", "P2-4": "覆盖",
    "P2-5": "覆盖", "P2-6": "覆盖", "P2-7": "部分覆盖", "P2-8": "覆盖",
    "P3-1": "覆盖", "P3-2": "覆盖", "P3-3": "覆盖", "P3-4": "覆盖",
    "P3-5": "覆盖", "P3-6": "未覆盖（待确认）",
}

OPEN_ITEMS = [
    ("外出打卡要素缺失", "TC-ATT-022",
     "文档只说明外出打卡需上传位置定位与拍照，未定义缺任一项是否阻断提交。"),
    ("考勤机引用后超限", "TC-ATT-017",
     "文档明确超限设备不可被选中，但未定义「规则已引用该设备后人数达上限」时打卡是否被拒。"),
    ("校准的部分时段匹配", "TC-ATT-038",
     "文档两例均以「审批是否匹配」为前提，未定义审批仅覆盖部分时段（如早退 2 小时仅请 1 小时假）时的校准粒度。"),
    ("未分配管辖范围的分级管理员", "TC-ATT-043",
     "文档未定义新建分级管理员未分配范围时的默认权限，属权限边界风险，需产品确认。"),
    ("打卡范围的配置取值区间", "TC-ATT-014",
     "文档给出「100至300米」有效区间，但未说明后台是否允许配置区间外的半径值。"),
    ("同一时段多类型审批并存", "TC-ATT-037",
     "文档列出五种审批模板可参与校准，未定义同一时段存在多张不同类型通过单据时的匹配优先级。"),
]


def fmt_cell(text):
    return str(text).replace("\n", "<br>").replace("|", "／")


# ---------------------------------------------------------------- markdown
def build_markdown():
    total = len(CASES)
    by_fp = OrderedDict()
    for c in CASES:
        by_fp.setdefault(c["fp"], []).append(c)

    sev_count = Counter(c["sev"] for c in CASES)
    method_count = Counter()
    tag_count = Counter()
    for c in CASES:
        for m in c["methods"].split(","):
            method_count[m.strip()] += 1
        for t in c["tags"].split(","):
            tag_count[t.strip()] += 1

    point_cases = {p: [] for p in POINTS}
    for c in CASES:
        for p in c["points"]:
            point_cases[p].append(c["id"])

    L = []
    w = L.append

    w("# 打卡管理 — 测试用例文档")
    w("")
    w(f"> 版本：v1.0 ｜ 编写日期：{TODAY} ｜ 用例总数：{total}")
    w("> 依据文档：`wiki/attendance-management/` 下 3 篇编译稿（打卡规则配置、特殊打卡规则与审批校准、分级管理权限与打卡数据导出）")
    w("> 需求原文：`raw/attendance-management/wecom-attendance-management.md`（企业微信《打卡管理》帮助文档）")
    w("")
    w("## 1. 文档说明")
    w("")
    w("### 1.1 范围")
    w("")
    w("本套用例覆盖打卡模块的管理侧配置与考勤侧判定，共 15 个功能点：规则类型与统计口径、打卡人员与白名单、打卡方式、"
      "打卡地点与范围、上下班时间与容差、外出打卡、拍照打卡、补卡申请、特殊日期打卡、打卡提醒、应用可见范围、"
      "审批关联与智能校准、分级管理权限、记录查看与导出、API 导出数据。")
    w("")
    w("### 1.2 依据与边界")
    w("")
    w("| 项 | 说明 |")
    w("| --- | --- |")
    w("| 唯一需求来源 | `wiki/attendance-management/*.md` 三篇文档 |")
    w("| 事实约束 | 用例中的配置入口、开关名称、数值、示例数据均取自文档原文，未引入外部经验值 |")
    w("| 文档未定义处 | 预期结果标注「待确认」，不臆造行为，统一登记在第 4.5 节待确认清单 |")
    w("| 不在范围 | 企业微信其他模块（通讯录、会议、日程）、终端系统兼容性、性能与容量压测 |")
    w("")
    w("### 1.3 严重等级定义（5 档）")
    w("")
    w("| 等级 | 名称 | 判定标准 | 用例数 |")
    w("| --- | --- | --- | --- |")
    for k in ("S1", "S2", "S3", "S4", "S5"):
        name, desc = SEVERITY[k]
        w(f"| {k} | {name} | {desc} | {sev_count.get(k, 0)} |")
    w("")
    w("### 1.4 设计方法与标签")
    w("")
    w("| 设计方法 | 使用次数 | 适用说明 |")
    w("| --- | --- | --- |")
    method_note = {
        "场景法": "按业务主线串联配置与打卡动作",
        "判定表": "多条件与开关组合下的判定规则",
        "边界值分析": "容差分钟数、范围半径、设备人数上限等阈值",
        "等价类划分": "部门/成员、范围内/外、状态分类",
        "错误推测": "越权、非常规路径、文档未定义边界",
        "状态迁移": "补卡审批前后的状态流转",
        "流程分析": "U 盘导入、权限分配、API 集成等多步链路",
        "矩阵分析": "多规则类型的横向对比",
    }
    for m, n in method_count.most_common():
        w(f"| {m} | {n} | {method_note.get(m, '')} |")
    w("")
    w("**关键标签**：" + "、".join(f"{t}({n})" for t, n in tag_count.most_common()))
    w("")
    w("## 2. 用例总览")
    w("")
    w("| 用例ID | 功能点 | 测试点 | 测试用例 | 等级 | 设计方法 | 标签 |")
    w("| --- | --- | --- | --- | --- | --- | --- |")
    for fp, cases in by_fp.items():
        for c in cases:
            w(f"| {c['id']} | {fmt_cell(FP[fp])} | {fmt_cell(c['point'])} | {fmt_cell(c['title'])} "
              f"| {c['sev']} | {fmt_cell(c['methods'])} | {fmt_cell(c['tags'])} |")
    w("")
    w("## 3. 用例明细")
    w("")
    for fp, cases in by_fp.items():
        w(f"### {fp} {FP[fp]}")
        w("")
        w(f"对应文档：`wiki/attendance-management/{FP_SECTION[fp][0]}` — {FP_SECTION[fp][1]}")
        w("")
        for c in cases:
            name = SEVERITY[c["sev"]][0]
            w(f"#### {c['id']} {c['title']}")
            w("")
            w("| 项 | 内容 |")
            w("| --- | --- |")
            w(f"| 功能点 | {fmt_cell(FP[c['fp']])} |")
            w(f"| 测试场景 | {fmt_cell(c['scene'])} |")
            w(f"| 测试点 | {fmt_cell(c['point'])} |")
            w(f"| 测试用例 | {fmt_cell(c['title'])} |")
            w(f"| 设计方法 | {fmt_cell(c['methods'])} |")
            w(f"| 关键标签 | {fmt_cell(c['tags'])} |")
            w(f"| 严重等级 | {c['sev']} {name} |")
            w(f"| 前置条件 | {fmt_cell(c['pre'])} |")
            w("")
            w("**测试步骤**")
            w("")
            for i, s in enumerate(c["steps"], 1):
                w(f"{i}. {s}")
            w("")
            w("**预期结果（测试结果）**")
            w("")
            for e in c["exp"]:
                w(f"- {e}")
            w("")
            w("**原始文档上下文**")
            w("")
            w(f"> {c['ctx']}")
            w(">")
            w(f"> —— `wiki/attendance-management/{c['src']}`")
            w("")
        w("")

    w("## 4. 覆盖度统计报告")
    w("")
    w("### 4.1 功能点覆盖")
    w("")
    w("| 功能点 | 名称 | 用例数 | S1 | S2 | S3 | S4 | S5 | 判定点数 |")
    w("| --- | --- | --- | --- | --- | --- | --- | --- | --- |")
    for fp, cases in by_fp.items():
        sc = Counter(c["sev"] for c in cases)
        pts = {p for c in cases for p in c["points"]}
        w(f"| {fp} | {FP[fp]} | {len(cases)} | {sc.get('S1', 0)} | {sc.get('S2', 0)} "
          f"| {sc.get('S3', 0)} | {sc.get('S4', 0)} | {sc.get('S5', 0)} | {len(pts)} |")
    w(f"| **合计** | — | **{total}** | **{sev_count.get('S1', 0)}** | **{sev_count.get('S2', 0)}** "
      f"| **{sev_count.get('S3', 0)}** | **{sev_count.get('S4', 0)}** | **{sev_count.get('S5', 0)}** | **{len(POINTS)}** |")
    w("")

    w("### 4.2 文档「可测试的关键判定点」覆盖矩阵")
    w("")
    w("| 判定点 | 内容 | 覆盖用例 | 状态 |")
    w("| --- | --- | --- | --- |")
    for p, text in POINTS.items():
        ids = "、".join(point_cases[p]) or "—"
        w(f"| {p} | {text} | {ids} | {POINT_STATUS[p]} |")
    w("")
    covered = sum(1 for p in POINTS if POINT_STATUS[p].startswith("覆盖"))
    partial = sum(1 for p in POINTS if POINT_STATUS[p] == "部分覆盖")
    pending = sum(1 for p in POINTS if POINT_STATUS[p].startswith("未覆盖"))
    cov_rate = (covered + 0.5 * partial) / len(POINTS) * 100
    w(f"- 判定点总数：{len(POINTS)}；完全覆盖 {covered} 条，部分覆盖 {partial} 条，待确认 {pending} 条")
    w(f"- **判定点覆盖率 = ({covered} + 0.5×{partial}) / {len(POINTS)} = {cov_rate:.1f}%**")
    w("")
    w("### 4.3 文档章节覆盖")
    w("")
    w("| 文档 | 章节 | 对应用例 |")
    w("| --- | --- | --- |")
    for fp, (src, sec) in FP_SECTION.items():
        ids = "、".join(c["id"] for c in by_fp[fp])
        w(f"| {src} | {sec} | {ids} |")
    w("")
    w("### 4.4 分布统计")
    w("")
    w("| 维度 | 分布 |")
    w("| --- | --- |")
    w("| 严重等级 | " + "、".join(f"{k} {SEVERITY[k][0]} {sev_count.get(k, 0)} 条" for k in ("S1", "S2", "S3", "S4", "S5")) + " |")
    w("| 设计方法 | " + "、".join(f"{m} {n}" for m, n in method_count.most_common()) + " |")
    w("| 关键标签 | " + "、".join(f"{t} {n}" for t, n in tag_count.most_common()) + " |")
    w("| 覆盖方式 | 正向 " + str(tag_count.get("正向", 0)) + " 条、异常 " + str(tag_count.get("异常", 0))
      + " 条、边界 " + str(tag_count.get("边界", 0)) + " 条、权限 " + str(tag_count.get("权限", 0)) + " 条 |")
    w("")
    w("### 4.5 未覆盖与待确认清单")
    w("")
    w("| # | 待确认项 | 关联用例 | 文档缺口说明 |")
    w("| --- | --- | --- | --- |")
    for i, (name, cid, note) in enumerate(OPEN_ITEMS, 1):
        w(f"| {i} | {name} | {cid} | {note} |")
    w("")
    w("### 4.6 结论")
    w("")
    w(f"1. 三篇文档的 15 个业务功能点均已建立用例，文档明确写出的 20 条可测试判定点中，{covered} 条完全覆盖、"
      f"{partial} 条部分覆盖，判定点覆盖率 {cov_rate:.1f}%。")
    w("2. 用例以正向主链路 + 异常/边界成对覆盖：正向 "
      f"{tag_count.get('正向', 0)} 条、异常 {tag_count.get('异常', 0)} 条、边界 {tag_count.get('边界', 0)} 条，"
      "避免只有正向用例导致缺陷漏出。")
    s1_ids = [c["id"] for c in CASES if c["sev"] == "S1"]
    w(f"3. 高风险面集中在四处：自由上下班的统计口径、手机打卡的「或」判定、审批校准的数据准确性、分级管理员的越权访问，"
      f"共 {len(s1_ids)} 条 S1 用例（{'、'.join(s1_ids)}），建议列入冒烟与回归必跑集。")
    w(f"4. 共有 {len(OPEN_ITEMS)} 项文档未定义的行为，已登记为待确认项；确认结论前，相关用例的预期结果不作为缺陷判定依据。")
    w("")
    return "\n".join(L)


# ---------------------------------------------------------------- excel
def build_excel(path, md_text):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "测试用例"

    headers = ["用例ID", "功能点", "功能点名称", "测试场景", "测试点", "测试用例",
               "设计方法", "关键标签", "严重等级", "严重等级名称", "前置条件",
               "测试步骤", "预期结果（测试结果）", "原始文档上下文", "依据文档"]
    ws.append(headers)

    head_fill = PatternFill("solid", fgColor="1F4E79")
    head_font = Font(color="FFFFFF", bold=True, size=10)
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for c in ws[1]:
        c.fill = head_fill
        c.font = head_font
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = border

    sev_fill = {"S1": "FFC7CE", "S2": "FFEB9C", "S3": "DDEBF7", "S4": "E2EFDA", "S5": "F2F2F2"}

    for c in CASES:
        steps = "\n".join(f"{i}. {s}" for i, s in enumerate(c["steps"], 1))
        exp = "\n".join(f"• {e}" for e in c["exp"])
        ws.append([
            c["id"], c["fp"], FP[c["fp"]], c["scene"], c["point"], c["title"],
            c["methods"], c["tags"], c["sev"], SEVERITY[c["sev"]][0], c["pre"],
            steps, exp, c["ctx"], f"wiki/attendance-management/{c['src']}",
        ])
        row = ws.max_row
        for cell in ws[row]:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = border
            cell.font = Font(size=10)
        ws.cell(row=row, column=9).fill = PatternFill("solid", fgColor=sev_fill[c["sev"]])
        ws.cell(row=row, column=1).font = Font(size=10, bold=True)

    widths = [12, 8, 18, 14, 30, 34, 16, 14, 9, 9, 26, 46, 40, 46, 36]
    for i, wd in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = wd
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:O{ws.max_row}"

    # 覆盖度 sheet
    ws2 = wb.create_sheet("覆盖度统计")
    ws2.append(["一、功能点覆盖"])
    ws2.append(["功能点", "名称", "用例数", "S1", "S2", "S3", "S4", "S5", "判定点数", "对应用例"])
    by_fp = OrderedDict()
    for c in CASES:
        by_fp.setdefault(c["fp"], []).append(c)
    for fp, cases in by_fp.items():
        sc = Counter(c["sev"] for c in cases)
        ids = "、".join(c["id"].replace("TC-ATT-", "") for c in cases)
        ws2.append([fp, FP[fp], len(cases), sc.get("S1", 0), sc.get("S2", 0), sc.get("S3", 0),
                    sc.get("S4", 0), sc.get("S5", 0),
                    len({p for c in cases for p in c["points"]}), ids])
    ws2.append([])
    ws2.append(["二、判定点覆盖矩阵"])
    ws2.append(["判定点", "内容", "覆盖用例", "状态"])
    point_cases = {p: [] for p in POINTS}
    for c in CASES:
        for p in c["points"]:
            point_cases[p].append(c["id"])
    for p, text in POINTS.items():
        ws2.append([p, text, "、".join(point_cases[p]) or "—", POINT_STATUS[p]])
    covered = sum(1 for p in POINTS if POINT_STATUS[p].startswith("覆盖"))
    partial = sum(1 for p in POINTS if POINT_STATUS[p] == "部分覆盖")
    cov_rate = (covered + 0.5 * partial) / len(POINTS) * 100
    ws2.append([])
    ws2.append(["判定点覆盖率", f"({covered} + 0.5×{partial}) / {len(POINTS)} = {cov_rate:.1f}%"])
    ws2.append([])
    ws2.append(["三、待确认项"])
    ws2.append(["#", "待确认项", "关联用例", "文档缺口说明"])
    for i, (name, cid, note) in enumerate(OPEN_ITEMS, 1):
        ws2.append([i, name, cid, note])

    sev_count = Counter(c["sev"] for c in CASES)
    method_count = Counter()
    tag_count = Counter()
    for c in CASES:
        for m in c["methods"].split(","):
            method_count[m.strip()] += 1
        for t in c["tags"].split(","):
            tag_count[t.strip()] += 1
    ws2.append([])
    ws2.append(["四、分布统计"])
    ws2.append(["维度", "分布"])
    ws2.append(["严重等级", "、".join(f"{k} {SEVERITY[k][0]} {sev_count.get(k, 0)} 条" for k in ("S1", "S2", "S3", "S4", "S5"))])
    ws2.append(["设计方法", "、".join(f"{m} {n}" for m, n in method_count.most_common())])
    ws2.append(["关键标签", "、".join(f"{t} {n}" for t, n in tag_count.most_common())])
    ws2.append(["用例总数", len(CASES)])

    head_rows = {"一、功能点覆盖", "功能点", "二、判定点覆盖矩阵", "判定点", "三、待确认项", "#", "四、分布统计", "维度"}
    for row in ws2.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        first = row[0].value
        if first in head_rows:
            for cell in row:
                if cell.value:
                    cell.font = Font(bold=True, color="FFFFFF")
                    cell.fill = head_fill
    for col, wd in zip("ABCDEFGHIJ", [14, 34, 10, 6, 6, 6, 6, 6, 10, 40]):
        ws2.column_dimensions[col].width = wd
    ws2.column_dimensions["B"].width = 46

    wb.save(path)


def main():
    md = build_markdown()
    with open("attendance-management-test-cases.md", "w", encoding="utf-8") as f:
        f.write(md + "\n")
    build_excel("attendance-management-test-cases.xlsx", md)
    print(f"cases={len(CASES)} md={len(md)} chars")


if __name__ == "__main__":
    main()
